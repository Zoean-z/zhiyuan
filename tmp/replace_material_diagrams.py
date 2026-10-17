from pathlib import Path
from docx import Document
from docx.shared import Inches,Emu

base=Path(r'D:\智愿ai改版\智愿AI报考平台_参赛资料_整改版')
diagrams=Path('tmp/material_diagrams').resolve()
mapping={
 '02_概要设计说明书.docx':{
  '图5-1 核心智能推荐流程图':'score-flow.png',
  '图5-2 AI Agent工作流程图':'agent-flow.png',
 },
 '03_详细设计说明书.docx':{
  '图4-1 双维度概率评分原理图':'score-flow.png',
  '图5-3 智能推荐业务流程图':'score-flow.png',
  '图5-4 AI Agent工作流程图':'agent-flow.png',
 }}
for name,items in mapping.items():
 p=base/name;d=Document(p)
 for caption,img in items.items():
  ps=list(d.paragraphs)
  index=next((i for i,q in enumerate(ps) if q.text.strip().startswith(caption)),None)
  if index is None:raise RuntimeError(caption)
  q=next((q for q in reversed(ps[:index]) if q._p.xpath('.//w:drawing')),None)
  if q is None:raise RuntimeError('image missing '+caption)
  ex=q._p.xpath('.//wp:extent')
  width=min(int(ex[0].get('cx')) if ex else Inches(6.2),Inches(6.2))
  for drawing in list(q._p.xpath('.//w:drawing')):drawing.getparent().remove(drawing)
  q.add_run().add_picture(str(diagrams/img),width=Emu(width))
 d.save(p)
 print('diagram replacements',name,len(items))

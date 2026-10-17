from pathlib import Path
from docx import Document
from docx.shared import Emu, Inches
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

OUT=Path(r'D:\智愿ai改版\智愿AI报考平台_参赛资料_整改版')
SHOTS=OUT/'截图证据'

def replace(doc, caption, shot):
    ps=list(doc.paragraphs)
    cap=next((i for i,p in enumerate(ps) if p.text.strip().startswith(caption)),None)
    if cap is None: raise RuntimeError(caption)
    p=next((p for p in reversed(ps[:cap]) if p._p.xpath('.//w:drawing')),None)
    if p is None: raise RuntimeError('image before '+caption)
    ex=p._p.xpath('.//wp:extent')
    width=min(int(ex[0].get('cx')) if ex else Inches(6.3),Inches(6.3))
    for drawing in list(p._p.xpath('.//w:drawing')):
        drawing.getparent().remove(drawing)
    p.add_run().add_picture(str(SHOTS/shot),width=Emu(width))

def add_after(doc, caption, shot, label):
    p=next(p for p in doc.paragraphs if p.text.strip().startswith(caption))
    def next_p(previous):
        el=OxmlElement('w:p');previous._p.addnext(el)
        return Paragraph(el,previous._parent)
    img=next_p(p)
    img.add_run().add_picture(str(SHOTS/shot),width=Inches(6.3))
    cap=next_p(img);cap.style=p.style;cap.add_run(label)

mapping={
 '05_软件界面设计书.docx':{
   '图3-4 系统首页':'home.png',
   '图3-6 查大学':'schools.png',
   '图3-7 院校详情':'school-detail.png',
   '图3-8 查专业':'majors.png',
   '图3-9 专业详情':'major-detail.png',
   '图3-12 志愿填报':'sheet.png',
   '图3-13 智能选大学':'choose.png',
   '图3-14 院校排行':'rank.png',
   '图3-18 AI对话':'agent.png',
   '图3-19 我的志愿方案':'plans.png',
   '图3-20 历史记录':'history.png',
   '图4-1 管理后台':'admin.png',
   '图4-2 用户管理':'admin.png',
 },
 '06_用户操作手册.docx':{
   '图2-4 系统首页':'home.png',
   '图3-1 查大学':'schools.png',
   '图3-2 院校详情':'school-detail.png',
   '图3-3 院校排行':'rank.png',
   '图4-1 查专业':'majors.png',
   '图4-2 专业详情':'major-detail.png',
   '图5-1 智能选大学':'choose.png',
   '图6-1 AI对话':'agent.png',
   '图7-1 志愿填报':'sheet.png',
   '图8-1 我的志愿方案':'plans.png',
   '图8-2 历史记录':'history.png',
 }}
for name,replacements in mapping.items():
    d=Document(OUT/name)
    for caption,shot in replacements.items(): replace(d,caption,shot)
    if name.startswith('05'):
        add_after(d,'图3-7 院校详情','school-detail-normal.png','图3-7-1 院校详情正常概率态')
        add_after(d,'图3-12 志愿填报','sheet-builder.png','图3-12-1 志愿填报器双页签')
        add_after(d,'图3-14 院校排行','rank-medical.png','图3-14-1 医药类筛选结果')
        add_after(d,'AI 管理工作台通过','admin-ai.png','图4-7 AI 管理工作台')
    else:
        add_after(d,'图7-1 志愿填报','sheet-builder.png','图7-1-1 志愿填报器')
    d.save(OUT/name)
    print('updated screenshots',name,len(replacements))

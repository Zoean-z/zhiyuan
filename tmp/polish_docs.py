from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

base=Path(r'D:\智愿ai改版\智愿AI报考平台_参赛资料_整改版')
p=base/'03_详细设计说明书.docx'
d=Document(p)
dups=[x for x in d.paragraphs if x.text.startswith('位次概率按三段计算')]
if len(dups)!=4:raise RuntimeError('unexpected score paragraphs')
el=dups[-1]._element;el.getparent().remove(el)
d.save(p)
for name,labels in {
 '05_软件界面设计书.docx':['图3-7-1','图3-12-1','图3-14-1','图4-7'],
 '06_用户操作手册.docx':['图7-1-1'],
}.items():
 p=base/name;d=Document(p)
 for x in d.paragraphs:
  if any(x.text.strip().startswith(v) for v in labels):
   x.alignment=WD_ALIGN_PARAGRAPH.CENTER
   for run in x.runs:run.bold=True
 d.save(p)
 print('polished',name)

from pathlib import Path
from docx import Document
import json
import sys

base = Path(r'D:\智愿ai改版\智愿AI报考平台_参赛资料')
for name in sys.argv[1:]:
    doc = Document(base / name)
    out = []
    for i, p in enumerate(doc.paragraphs):
        s = p.text.strip()
        if s:
            out.append({'kind': 'p', 'i': i, 'style': p.style.name, 'text': s[:500]})
    for i, t in enumerate(doc.tables):
        out.append({'kind': 'table', 'i': i, 'rows': len(t.rows), 'cols': len(t.columns), 'head': [[c.text[:100] for c in row.cells] for row in t.rows[:3]]})
    print(json.dumps({'file': name, 'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables), 'images': len(doc.inline_shapes), 'items': out}, ensure_ascii=False))

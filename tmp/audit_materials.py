from docx import Document
from pathlib import Path
import re

out=Path(r'D:\智愿ai改版\智愿AI报考平台_参赛资料_整改版')
pat=re.compile(r'80\s*所|100\s*个专业|每组最多\s*5|最多返回\s*5|截取\s*Top\s*5|共\s*15\s*所|106\s*(?:个|项|测试)|9\s*个(?:预定义|白名单|业务)?工具|双层意图决策|本地正则优先|前端.{0,12}轮询|没有AI配置|无AI配置|AI配置管理未实现|PK对比|Ctrl\+K|NProgress|哆啦A梦|面包屑|三[个步]步骤|第三步|修改密码',re.I)
for p in sorted(out.glob('0[1-6]*.docx')):
 d=Document(p)
 rows=[]
 for i,x in enumerate(d.paragraphs):
  if pat.search(x.text): rows.append(('P',i,x.text[:230]))
 for i,t in enumerate(d.tables):
  for j,row in enumerate(t.rows):
   s=' | '.join(c.text for c in row.cells)
   if pat.search(s): rows.append(('T',f'{i}:{j}',s[:230]))
 print('\n'+p.name,'hits',len(rows))
 for v in rows[:150]:print(*v)

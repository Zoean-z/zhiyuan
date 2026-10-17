from docx import Document
import sys
from pathlib import Path

base=Path(r'D:\智愿ai改版\智愿AI报考平台_参赛资料')
for arg in sys.argv[1:]:
 d=Document(base/arg)
 print(arg)
 for i,p in enumerate(d.paragraphs):
  blips=p._p.xpath('.//a:blip')
  if not blips:continue
  prev=next((x.text.strip() for x in reversed(d.paragraphs[:i]) if x.text.strip()),'')
  nxt=next((x.text.strip() for x in d.paragraphs[i+1:] if x.text.strip()),'')
  print(i,[(b.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'),d.part.related_parts[b.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')].partname) for b in blips],prev[:60],'=>',nxt[:60])

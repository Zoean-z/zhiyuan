from pathlib import Path
from docx import Document

p=Path(r'D:\智愿ai改版\智愿AI报考平台_参赛资料_整改版\04_数据库设计说明书.docx')
d=Document(p)
for x in d.paragraphs:
    if x.text.startswith('空数据卷首次部署由 MySQL initdb'):
        s=('空数据卷首次部署由 MySQL initdb 执行 sql/schema.sql、sql/data.sql、'
           'sql/upgrade-20260822-unique-keys.sql；DB_SCHEMA_INIT_MODE 默认 never。'
           'sql/data.sql 为 80 校轻量种子。需要复现 zhiyuan6 的公开全量数据时，'
           '叠加 docker-compose.full-data.yml，随后执行 sql/full-competition-data-20260923.sql，'
           '导入 1,902 校与七省录取事实。该快照不含用户、方案、对话和 AI 密钥。'
           '历史增量脚本另包括 rebuild-20260825-universities、upgrade-20260825-vocational-majors、'
           'extend-hunan-multi-year-20260828、init-competition-admission-cutoffs 和 import-score-rank-mapping。')
        if x.runs:
            x.runs[0].text=s
            for r in x.runs[1:]:r.text=''
        else:x.add_run(s)
        break
else:raise RuntimeError('database init paragraph not found')
d.save(p)
print('updated',p.name)

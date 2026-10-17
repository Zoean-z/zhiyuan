from pathlib import Path
from io import BytesIO
from docx import Document
from PIL import Image, ImageDraw, ImageFont, ImageOps

root=Path(r'D:\智愿ai改版\智愿AI报考平台_参赛资料_整改版')
out=Path('tmp/review_20260924_images')
out.mkdir(exist_ok=True)
targets={
 '03': ['图7-4','图7-5','图7-7','图7-9','图7-10','图7-11','图7-12','图7-13','图7-15','图7-16','图7-19','图7-20','图7-21'],
 '05': ['图3-5','图3-15','图3-16','图3-17','图4-3','图4-4','图4-5','图4-6'],
 '06': ['图2-5','图5-2','图5-3','图5-4','图9-1','图9-2'],
}
rel='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'
font_path=Path(r'C:\Windows\Fonts\msyh.ttc')
font=ImageFont.truetype(str(font_path),18) if font_path.exists() else ImageFont.load_default()
for prefix,captions in targets.items():
    source=next(root.glob(prefix+'*.docx'))
    doc=Document(source)
    images=[]
    for cap in captions:
        index=next((i for i,p in enumerate(doc.paragraphs) if p.text.strip().startswith(cap)),None)
        if index is None:
            print(prefix,cap,'MISSING CAPTION')
            continue
        p=next((p for p in reversed(doc.paragraphs[:index]) if p._p.xpath('.//a:blip')),None)
        if p is None:
            print(prefix,cap,'MISSING IMAGE')
            continue
        rid=p._p.xpath('.//a:blip')[0].get(rel)
        blob=doc.part.related_parts[rid].blob
        try: im=Image.open(BytesIO(blob)).convert('RGB')
        except Exception as error:
            print(prefix,cap,'UNREADABLE',error)
            continue
        path=out/(prefix+'_'+cap.replace('图','fig').replace('-','_')+'.png')
        im.save(path)
        images.append((cap,im,path))
        print(prefix,cap,im.size,path)
    if not images:continue
    tile_w,tile_h=420,285
    cols=3;rows=(len(images)+cols-1)//cols
    sheet=Image.new('RGB',(cols*tile_w,rows*tile_h),'white')
    draw=ImageDraw.Draw(sheet)
    for i,(cap,im,path) in enumerate(images):
        x=(i%cols)*tile_w;y=(i//cols)*tile_h
        thumb=ImageOps.contain(im,(tile_w-16,tile_h-42))
        sheet.paste(thumb,(x+(tile_w-thumb.width)//2,y+34))
        draw.text((x+8,y+6),cap,font=font,fill='#111111')
    sheet.save(out/(prefix+'_contact.png'))

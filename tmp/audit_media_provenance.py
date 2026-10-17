from pathlib import Path
from hashlib import sha256
from docx import Document

current=Path(r'D:\智愿ai改版\智愿AI报考平台_参赛资料_整改版')
original=Path(r'D:\智愿ai改版\智愿AI报考平台_参赛资料')
shots={sha256(p.read_bytes()).hexdigest():p.name for p in (current/'截图证据').glob('*.png')}
rel_ns='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'
for prefix in ('02','03','05','06'):
    file=next(current.glob(prefix+'*.docx'))
    old=Document(original/file.name)
    old_hashes={sha256(part.blob).hexdigest() for part in old.part.related_parts.values() if str(part.content_type).startswith('image/')}
    doc=Document(file)
    counts={}
    rows=[]
    for i,p in enumerate(doc.paragraphs):
        blips=p._p.xpath('.//a:blip')
        for b in blips:
            rid=b.get(rel_ns)
            if not rid:continue
            digest=sha256(doc.part.related_parts[rid].blob).hexdigest()
            source=('live_capture:'+shots[digest] if digest in shots else 'original_image' if digest in old_hashes else 'new_image')
            cap=next((q.text.strip() for q in doc.paragraphs[i+1:i+5] if q.text.strip()),'')
            counts[source]=counts.get(source,0)+1
            rows.append((i,source,cap[:75]))
    print('\n'+file.name,'image_status',counts)
    for row in rows:print(*row)
